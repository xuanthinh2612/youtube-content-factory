from pathlib import Path
from contextlib import redirect_stderr,redirect_stdout
import io,json,re,html
import bleach
import markdown as md
from docx import Document
from quality.validators.duration import clean_narration_text


CSS="""@page{size:A4;margin:20mm}
body{font-family:'Noto Sans','Noto Sans CJK JP',Arial,sans-serif;line-height:1.65;color:#111827}
h1,h2,h3{line-height:1.25}
table{border-collapse:collapse;width:100%}
td,th{border:1px solid #ddd;padding:6px}
p{margin:.65em 0 1em}
"""

ALLOWED_TAGS=set(bleach.sanitizer.ALLOWED_TAGS)|{
    "p","h1","h2","h3","h4","h5","h6","pre","code","blockquote",
    "ul","ol","li","hr","br","table","thead","tbody","tr","th","td"
}
ALLOWED_ATTRS={"a":["href","title","rel"],"code":["class"],"pre":["class"]}

_PDF_HTML=None
_PDF_IMPORT_ERROR=None


def _pdf_renderer():
    """Load the optional native PDF stack once and keep failures non-fatal."""
    global _PDF_HTML,_PDF_IMPORT_ERROR
    if _PDF_HTML is not None:
        return _PDF_HTML
    if _PDF_IMPORT_ERROR is not None:
        raise RuntimeError(_PDF_IMPORT_ERROR)
    try:
        # Some WeasyPrint builds print a long native-library diagnostic before
        # raising.  The concise exception is surfaced in our export warning.
        with redirect_stdout(io.StringIO()),redirect_stderr(io.StringIO()):
            from weasyprint import HTML
        _PDF_HTML=HTML
        return HTML
    except Exception as exc:
        _PDF_IMPORT_ERROR=f"{type(exc).__name__}: {exc}"
        raise RuntimeError(_PDF_IMPORT_ERROR) from exc


def _has_h1(text):
    for line in text.splitlines():
        s=line.strip()
        if not s:
            continue
        return s.startswith("# ")
    return False


def export_all(folder:Path,stem,title,language,text,metadata=None,tts_ready=False,include_title=True):
    folder.mkdir(parents=True,exist_ok=True)
    metadata=metadata or {}
    paths={};warnings=[]

    # Essential, plain-text artifacts. If these fail, the export truly failed.
    mdp=folder/f"{stem}.md"
    mdp.write_text(text,encoding="utf-8")
    paths["md"]=str(mdp)

    narration_text=clean_narration_text(text) if tts_ready else ""
    if tts_ready:
        tp=folder/f"{stem}.txt"
        tp.write_text(narration_text,encoding="utf-8")
        paths["txt"]=str(tp)

    jp=folder/f"{stem}.json"
    jp.write_text(json.dumps({
        "title":title,"language":language,"content":text,"narration_text":narration_text,"metadata":metadata
    },ensure_ascii=False,indent=2),encoding="utf-8")
    paths["json"]=str(jp)

    # Rich formats are best-effort. A font/WeasyPrint/docx issue must not turn a
    # completed content workflow into Project stopped after the canonical text is safe.
    try:
        doc=Document()
        if include_title and title and not _has_h1(text):
            doc.add_heading(title,0)
        for line in text.splitlines():
            s=line.strip()
            if not s:
                continue
            if s.startswith("### "):doc.add_heading(s[4:],3)
            elif s.startswith("## "):doc.add_heading(s[3:],2)
            elif s.startswith("# "):doc.add_heading(s[2:],1)
            elif s.startswith("- "):doc.add_paragraph(s[2:],style="List Bullet")
            else:doc.add_paragraph(re.sub(r"\*\*(.*?)\*\*",r"\1",s))
        dp=folder/f"{stem}.docx"
        doc.save(dp)
        paths["docx"]=str(dp)
    except Exception as exc:
        warnings.append(f"DOCX export failed: {type(exc).__name__}: {exc}")

    try:
        # WeasyPrint depends on native Pango/GLib libraries which are commonly
        # absent on a plain Windows Python installation.  PDF is an optional rich
        # export, so importing it at module load time must not prevent the API,
        # workers, or essential Markdown/JSON/TXT exports from starting.
        HTML=_pdf_renderer()
        rendered=md.markdown(text,extensions=["tables","fenced_code","sane_lists"])
        rendered=bleach.clean(
            rendered,tags=ALLOWED_TAGS,attributes=ALLOWED_ATTRS,
            protocols={"http","https","mailto"},strip=True
        )
        title_html="" if _has_h1(text) or not title or not include_title else f"<h1>{html.escape(title)}</h1>"
        hp=f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{title_html}{rendered}</body></html>"
        pp=folder/f"{stem}.pdf"
        HTML(string=hp,base_url=str(folder)).write_pdf(pp)
        paths["pdf"]=str(pp)
    except Exception as exc:
        warnings.append(f"PDF export failed: {type(exc).__name__}: {exc}")

    if warnings:
        paths["_warnings"]=warnings
    return paths
